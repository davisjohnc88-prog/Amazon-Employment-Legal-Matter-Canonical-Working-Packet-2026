# Reads every document back after make_package.py and checks the folders. No AI involved.
# Usage: python3 verify_package.py "<out_dir>/Workplace Psych"
import os, re, sys
from docx import Document

ROOT = sys.argv[1]
block = [r"\bamazon\b", r"\baws\b", r"whole foods", r"wow project", r"fulfil+ment cent(er|re)", r"\bprime day\b", r"\bpeak season\b"]
priv = "/mnt/project-files/series/private-blocklist.txt"
if os.path.exists(priv):
    block += [re.escape(l.strip()) for l in open(priv) if l.strip() and not l.startswith("#")]
from spellchecker import SpellChecker
sp = SpellChecker()
ok = set()
for fn in ("../trees/wordlist.txt", "../build/wordlist.txt"):
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), fn)
    ok |= {w.strip().lower() for w in open(p) if w.strip()}
ok |= {"pdf", "png", "docx", "linkedin", "instagram", "substack", "alt", "darvo", "pip", "ig", "fb"}
problems, docs, folders, seen = [], 0, 0, set()


def text(path):
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out += [c.text for c in r.cells]
    return "\n".join(out)


for dirpath, dirs, files in os.walk(ROOT):
    for f in files:
        if f.endswith(".docx"):
            docs += 1
            t = text(os.path.join(dirpath, f))
            rel = os.path.relpath(os.path.join(dirpath, f), ROOT)
            for b in block:
                m = re.search(b, t, re.I)
                if m:
                    problems.append(f"employer/private term in {rel}: {m.group()!r}")
            if f == "Post wording.docx":
                body = re.sub(r"https?://\S+|doi\.org/\S+|#\w+", " ", t.split("Whose words are these")[0])
                for w in re.findall(r"[^\W\d_][^\W\d_'’]*(?:['’][^\W\d_]+)?", body):
                    lw = w.lower().replace("’", "'")
                    lw = lw[:-2] if lw.endswith("'s") else lw
                    if len(lw) > 2 and not w[0].isupper() and lw not in ok and lw not in seen and sp.unknown([lw]):
                        seen.add(lw)
                        problems.append(f"spelling? {rel}: {w!r}")
    if "Post wording.docx" in files:
        folders += 1
        rel = os.path.relpath(dirpath, ROOT)
        t = text(os.path.join(dirpath, "Post wording.docx"))
        pngs = [f for f in files if f.endswith(".png")]
        alts = re.findall(r"^(?:Slide \d+|Image): ", t, re.M)
        if len(pngs) != len(alts):
            problems.append(f"{rel}: {len(pngs)} images but {len(alts)} alt texts")
        if len(pngs) > 1 and "LinkedIn carousel.pdf" not in files:
            problems.append(f"{rel}: carousel without its PDF")
        cap = t.split("Caption (copy and paste)")[1].split("Alt text (paste")[0]
        n, lim = (int(x.replace(",", "")) for x in re.search(r"([\d,]+) of ([\d,]+) characters", cap).groups())
        if n > lim:
            problems.append(f"{rel}: caption too long ({n})")
        if re.search(r"crisis|suicid", cap, re.I) and "988" not in cap:
            problems.append(f"{rel}: crisis mention without 988")
        for f in pngs + [x for x in files if x.endswith(".pdf")]:
            if os.path.getsize(os.path.join(dirpath, f)) < 20000:
                problems.append(f"{rel}: {f} looks empty")

print(f"read back {docs} documents; {folders} post folders")
print("\n".join(problems) if problems else "no problems")
sys.exit(1 if problems else 0)
